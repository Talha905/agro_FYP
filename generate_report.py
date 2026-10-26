import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()

    # Page Setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Color Palette Constants
    HEX_PRIMARY_NAVY = "1B365D"
    HEX_SECONDARY_BLUE = "2C4D75"
    HEX_LIGHT_BG = "F4F6F8"
    HEX_BORDER = "D3D3D3"
    COLOR_PRIMARY_NAVY = RGBColor(0x1B, 0x36, 0x5D)
    COLOR_SECONDARY_BLUE = RGBColor(0x2C, 0x4D, 0x75)
    COLOR_TEXT_MAIN = RGBColor(0x22, 0x22, 0x22)
    COLOR_MUTED = RGBColor(0x66, 0x66, 0x66)

    # Set Default Normal Style
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = COLOR_TEXT_MAIN

    # Helper Functions
    def add_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY_NAVY
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.italic = True
        run.font.color.rgb = COLOR_MUTED
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY_NAVY
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY_BLUE
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY_BLUE
        return p

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.bold = True
        p.add_run(text)
        return p

    def add_bullet(text, bold_prefix=None, level=0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.bold = True
        p.add_run(text)
        return p

    def set_cell_background(cell, hex_color):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def style_table(table, col_widths, col_alignments=None):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row_idx, row in enumerate(table.rows):
            # CantSplit
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            if row_idx == 0:
                trPr.append(OxmlElement('w:tblHeader'))

            for col_idx, cell in enumerate(row.cells):
                cell.width = col_widths[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                # Align text
                if col_alignments and col_idx < len(col_alignments):
                    for p in cell.paragraphs:
                        p.alignment = col_alignments[col_idx]
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)

                if row_idx == 0:
                    set_cell_background(cell, HEX_PRIMARY_NAVY)
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "FFFFFF")
                    else:
                        set_cell_background(cell, HEX_LIGHT_BG)
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.color.rgb = COLOR_TEXT_MAIN

    # --- DOCUMENT HEADER ---
    add_title("AgroSaathi – Updated GitHub FYP Progress Audit")
    add_subtitle("Branch Analysis, Individual Contribution Report & Integration Review")

    add_body("Project: AgroSaathi – Smart Farming Assistant")
    add_body("Repository: Talha905/agro_FYP")
    add_body("Branch Audited: feature-recommender (supersedes prior develop branch audit)")
    add_body("Original Audit Date: 21 August 2026")
    add_body("Previous Update: 27 August 2026")
    add_body("This Update: 1 October 2026")
    add_body("Assessment Role: Senior Software Engineer, FYP Evaluator and Project Supervisor")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 1 ---
    add_heading_1("1. Purpose of This Update")
    add_body(
        "The first progress report (21 Aug 2026) audited four disconnected feature-branch snapshots, and the second report "
        "(27 Aug 2026) audited their initial consolidation into the develop branch. This third report audits the feature-recommender branch "
        "as of 1 October 2026, incorporating all recent architectural refactoring, location resolution fixes, Firestore index optimization, "
        "community recommendation voting, and vision-assisted crop growth plan recalibration."
    )
    add_body(
        "This report re-evaluates the codebase against the original Requirements Traceability Matrix and restates module ownership using actual team member names:"
    )
    add_bullet(" — Dashboard + Crop Recommender", bold_prefix="• Person A → Shubham")
    add_bullet(" — Authentication + Business Helper / Bidding", bold_prefix="• Person B → Bhavesh")
    add_bullet(" — Crop Disease Detector", bold_prefix="• Person C → Kunal")
    add_bullet(" — Crop Growth Planner", bold_prefix="• Person D → Talha")

    # --- SECTION 2 ---
    add_heading_1("2. Headline Numbers: Then vs. Now")
    table2 = doc.add_table(rows=6, cols=4)
    col_widths2 = [Inches(2.2), Inches(1.3), Inches(1.5), Inches(1.5)]
    col_align2 = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]

    headers2 = ["Area", "Previous (27 Aug)", "Current (1 Oct)", "Change"]
    for i, h in enumerate(headers2):
        table2.cell(0, i).paragraphs[0].text = h

    data2 = [
        ["Overall functional development", "68%", "82%", "▲ +14 pts"],
        ["Mainline-integrated progress", "65%", "80%", "▲ +15 pts"],
        ["Integration readiness score", "72 / 100", "84 / 100", "▲ +12 pts"],
        ["Project stage", "Feature Development → Integration", "Advanced Integration & Refinement", "▲ Refined"],
        ["Final submission readiness", "Moderate to High", "High", "▲ High Readiness"]
    ]

    for row_idx, row_data in enumerate(data2, start=1):
        for col_idx, cell_value in enumerate(row_data):
            table2.cell(row_idx, col_idx).paragraphs[0].text = cell_value

    style_table(table2, col_widths2, col_align2)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 3 ---
    add_heading_1("3. What Changed in the Feature-Recommender Branch (Evidence-Based)")
    add_body(
        "All claims below were directly verified against source code in agrosaathi/lib/, backend/app.py, and firestore.rules on branch feature-recommender as of 1 October 2026."
    )

    # 3.1
    add_heading_2("3.1 Dashboard & Location Resolution — Live GPS Integration")
    add_body("Previously: The dashboard weather card and crop recommender defaulted silently to 'Pune, Maharashtra' regardless of user device location.")
    add_body("Now: A shared, singleton location service eliminates silent fallbacks and enforces device GPS primacy.")
    add_bullet("Shared Location State: AppLocationProvider.stateNotifier (agrosaathi/lib/services/app_location_provider.dart, L18-164) provides a single source of location state consumed synchronously by both WeatherWidget (agrosaathi/lib/widgets/weather_widget.dart, L60-202) and CropAdvisorScreen (agrosaathi/lib/screens/crop_advisor_screen.dart, L78-97).", bold_prefix="[Implemented & Verified] ")
    add_bullet("Live GPS Primacy & Fallback: Fresh GPS coordinates win over saved data. Last detected location is saved to SharedPreferences ('last_known_district', L75) strictly as an offline fallback.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Reverse Geocoding: LocationService.reverseGeocode() (agrosaathi/lib/services/location_service.dart, L134-225) primary geocoding via Flutter geocoding package (placemarkFromCoordinates) with Nominatim API fallback (User-Agent: AgroSaathiApp/1.0, 1s rate-limiting, 2-decimal coordinate caching). LocationService.extractBroadDistrict() (L69-82) strips administrative division names ('Konkan Division') to extract clean district names.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Explicit Location Error Handling: Explicit UI cards and action buttons handle LocationErrorType.serviceDisabled (button calls Geolocator.openLocationSettings()), LocationErrorType.permissionDenied (calls requestPermission() / openAppSettings()), and LocationErrorType.timeout (15s timeLimit retry button, L247-325).", bold_prefix="[Implemented & Verified] ")

    # 3.2
    add_heading_2("3.2 Crop Recommender & Community Presets — Real-Data & Voting")
    add_body("Previously: Hard-coded static soil baseline presets auto-filled input text fields, and no community feedback loop existed.")
    add_body("Now: Static presets are replaced by live environmental data and dynamic, vote-ranked community recommendations.")
    add_bullet("Live Weather & Soil Integration: Sourced dynamically from Open-Meteo REST API (https://api.open-meteo.com/v1/forecast via WeatherService.fetchWeatherByCoordinates()) and SoilGrids REST API (https://rest.isric.org/soilgrids/v2.0/properties/query via SoilGridsService.fetchSoilProperties()).", bold_prefix="[Implemented & Verified] ")
    add_bullet("Top 3 Comparative Preset Cards: CommunityPresetService.fetchPresetsForUser() (agrosaathi/lib/services/community_preset_service.dart, L112-149) ranks presets by net score (usefulVotes - notUsefulVotes). Displays top 3 cards with N, P, K, pH, and water specs alongside database-only vote counts (👍 {usefulVotes} 👎 {notUsefulVotes}). Tapping a card highlights it with a primary border and populates form fields.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Realistic Soil Seeds: Seed presets generated in getSeedPresets() (L78-107) define 3 variants per soil type (Standard Baseline, High-Yield Intensive, Organic/Low Input), labelled 'Typical for <soil type>' with 👍 0 👎 0.", bold_prefix="[Implemented & Verified] ")
    add_bullet("District-Level Aggregation: Presets are keyed by broad district via extractBroadDistrict() so votes accumulate across the district.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Post-Recommendation Feedback Card: RecommendationResultScreen (agrosaathi/lib/screens/recommendation_result_screen.dart, L148-203) presents an in-app card ('Was this recommendation useful for {district}?') that calls CommunityPresetService.recordRecommendationVote() to record votes transactionally.", bold_prefix="[Implemented & Verified] ")
    add_bullet("SoilGrids pH Indicator: Displays an explicit 'Soil pH (Fetched via SoilGrids)' badge when pH is prefilled from live SoilGrids API.", bold_prefix="[Implemented & Verified] ")

    # 3.3
    add_heading_2("3.3 Past History & Firestore Indexes — Query Optimization")
    add_body("Previously: Viewing past history threw a Cloud Firestore exception ([cloud_firestore/failed-precondition] query requires index).")
    add_body("Now: Resolved without requiring custom composite index deployment while maintaining cache resilience.")
    add_bullet("Single-Field Query Fix: CropRecommendationService.streamHistory() (agrosaathi/lib/services/crop_recommendation_service.dart, L380) updated to query .where('farmerId', isEqualTo: farmerId) only and sort client-side (records.sort((a, b) => b.createdAt.compareTo(a.createdAt))), eliminating runtime index dependencies.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Index Schema Config: firestore.indexes.json and firebase.json created to define composite index schemas for optional production deployment via firebase deploy --only firestore:indexes.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Cache Preservation: CropAdvisorScreen._buildHistoryTab() (agrosaathi/lib/screens/crop_advisor_screen.dart, L583-620) retains cached history records on stream errors to prevent UI list resets during network reconnects.", bold_prefix="[Implemented & Verified] ")

    # 3.4
    add_heading_2("3.4 Growth Planner Backend & Ollama Integration — qwen2.5:7b")
    add_body("Previously: The backend prompt contained hardcoded stage durations and fixed NPK numbers, biasing LLM outputs.")
    add_body("Now: The prompt is unconstrained and schema-enforced via Ollama structured outputs.")
    add_bullet("Unconstrained Prompt: GROWTH_PLAN_SYSTEM_PROMPT (backend/app.py, L670-695) rewritten to state required data types without fixed example numbers.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Structured Output & Schema Mode: Ollama API called with format param matching JSON schema (cropName, stages, fertilizerPlan). Validated against reference duration table backend/data/crop_growth_durations.json.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Fallback Estimated Flag: When Ollama API is unavailable or times out, backend returns static template with 'isEstimated': True and 'source': 'estimated_fallback'.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Sanitization & Logging: Input cropName sanitized to strip script tags. Ollama call diagnostics logged (model: qwen2.5:7b, host: http://localhost:11434, latency, HTTP status).", bold_prefix="[Implemented & Verified] ")

    # 3.5
    add_heading_2("3.5 Growth Planner Image Feature — Vision-Assisted Recalibration")
    add_body("New Feature: Farmers can capture or upload crop photo to refine growth schedule based on real field observations.")
    add_bullet("Photo Upload Button: GrowthPlanDetailScreen (agrosaathi/lib/screens/growth_plan_detail_screen.dart, L525-580) adds full-width '📷 Upload Crop Photo to Refine Plan' button inside harvest countdown card.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Backend Endpoint & Vision Model: /adapt-growth-plan endpoint (backend/app.py, L750-820) integrates Qwen 2.5 VL vision model (qwen2.5vl) for crop stage/condition observation alongside PlantVillage disease classifier.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Confidence Threshold & Stage Shift Cap: Enforces confidence threshold (<60% prompts retake) and caps schedule adjustments to ±30%. Interactive diff summary dialog presents schedule changes before applying.", bold_prefix="[Implemented & Verified; Pending On-Device Camera Field Verification] ")

    # 3.6
    add_heading_2("3.6 Security, Firestore Rules & Configuration")
    add_bullet("Firestore Rules File: firestore.rules authored in repository root defining security rules for users, community_crop_presets, recommendation_votes, recommendations, and growth_plans.", bold_prefix="[Implemented Locally; Pending Firebase CLI Deploy] ")
    add_bullet("Firebase CLI Link: firebase.json configured to reference firestore.rules and firestore.indexes.json.", bold_prefix="[Implemented & Verified] ")
    add_bullet("Deployment Status: Rules are present in local codebase but NOT YET DEPLOYED to production Firebase project farmerhelp-4d27c. Must execute firebase deploy --only firestore:rules.", bold_prefix="[Pending Deploy] ")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 4 ---
    add_heading_1("4. Updated Requirements Traceability (Deltas Only)")
    add_body("The table below lists requirement rows whose implementation status changed since the 27 August audit.")

    table4 = doc.add_table(rows=7, cols=5)
    col_widths4 = [Inches(0.6), Inches(1.8), Inches(1.0), Inches(1.4), Inches(1.7)]
    col_align4 = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]

    headers4 = ["ID", "Requirement", "Owner", "Previous Status", "Current Status"]
    for i, h in enumerate(headers4):
        table4.cell(0, i).paragraphs[0].text = h

    data4 = [
        ["R14", "Location & Weather Integration", "Shubham", "Hardcoded Pune default", "Implemented — Shared AppLocationProvider, Open-Meteo, SoilGrids, reverse geocoding, error states"],
        ["R15", "Crop Recommender Presets & Voting", "Shubham", "Static presets in widget", "Implemented — Top 3 vote-ranked presets, DB counts, district keying, in-app voting card"],
        ["R16", "Recommendation History & Querying", "Shubham", "Blocked by Firestore index", "Implemented — Single-field Firestore query, client-side sorting, cache retention"],
        ["R25", "AI Growth Plan Schema Generation", "Talha", "Gemini 404 model blocker", "Implemented — qwen2.5:7b structured output via Ollama, duration validation, fallback flag"],
        ["R26", "Growth Plan Vision Recalibration", "Talha", "Not scoped", "Implemented — Qwen 2.5 VL vision endpoint, confidence threshold, schedule diff summary"],
        ["R36", "Database Security Rules", "All", "Missing (0 rules files)", "Implemented Locally — firestore.rules & firebase.json authored; CLI deploy pending"]
    ]

    for row_idx, row_data in enumerate(data4, start=1):
        for col_idx, cell_value in enumerate(row_data):
            table4.cell(row_idx, col_idx).paragraphs[0].text = cell_value

    style_table(table4, col_widths4, col_align4)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 5 ---
    add_heading_1("5. Individual Contribution Summary (Updated)")
    add_body(
        "Note: In accordance with audit rules, team member contribution percentages are kept exactly as reported in the 27 August update. "
        "This update (1 October 2026) reflects primary engineering deltas in Dashboard/Recommender (Shubham) and Growth Planner/Backend (Talha)."
    )

    table5 = doc.add_table(rows=5, cols=5)
    col_widths5 = [Inches(1.2), Inches(2.2), Inches(1.1), Inches(1.1), Inches(1.1)]
    col_align5 = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]

    headers5 = ["Team Member", "Responsibility", "21 Aug", "27 Aug / 1 Oct", "Change Note"]
    for i, h in enumerate(headers5):
        table5.cell(0, i).paragraphs[0].text = h

    data5 = [
        ["Shubham", "Dashboard + Crop Recommender", "65%", "68%", "Primary contributor for Oct 1 location, Open-Meteo/SoilGrids & preset voting updates"],
        ["Bhavesh", "Authentication + Business Helper / Bidding", "35%", "68%", "Unchanged since Aug 27 audit (Marketplace module fully built out)"],
        ["Kunal", "Crop Disease Detector", "55%", "62%", "Unchanged since Aug 27 audit (PlantVillage MobileNetV2 pipeline documented)"],
        ["Talha", "Crop Growth Planner", "75%", "78%", "Primary contributor for Oct 1 qwen2.5:7b structured schema & vision recalibration"]
    ]

    for row_idx, row_data in enumerate(data5, start=1):
        for col_idx, cell_value in enumerate(row_data):
            table5.cell(row_idx, col_idx).paragraphs[0].text = cell_value

    style_table(table5, col_widths5, col_align5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 6 ---
    add_heading_1("6. Revised Module Scorecard")
    table6 = doc.add_table(rows=14, cols=4)
    col_widths6 = [Inches(2.4), Inches(1.2), Inches(1.2), Inches(1.7)]
    col_align6 = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]

    headers6 = ["Module", "27 Aug", "1 Oct", "Trend & Justification"]
    for i, h in enumerate(headers6):
        table6.cell(0, i).paragraphs[0].text = h

    data6 = [
        ["Authentication", "60%", "60%", "→ Unchanged"],
        ["Dashboard", "72%", "85%", "↗ ▲ +13 pts (Live GPS, shared provider, error states)"],
        ["Crop Recommender", "76%", "88%", "↗ ▲ +12 pts (Open-Meteo, SoilGrids, top 3 vote presets)"],
        ["Disease Detector", "62%", "62%", "→ Unchanged"],
        ["Growth Planner", "78%", "90%", "↗ ▲ +12 pts (qwen2.5:7b schema, vision recalibration)"],
        ["Bidding / Marketplace", "80%", "80%", "→ Unchanged"],
        ["Multilingual Support", "75%", "75%", "→ Unchanged (in-memory locale switcher intact)"],
        ["Firebase / Database", "62%", "80%", "↗ ▲ +18 pts (Single-field index fix, local rules authored)"],
        ["Security", "15%", "50%", "↗ ▲ +35 pts (firestore.rules & firebase.json authored)"],
        ["ML", "62%", "75%", "↗ ▲ +13 pts (Qwen 2.5 VL vision integration)"],
        ["Integration", "70%", "82%", "↗ ▲ +12 pts (Single feature-recommender branch)"],
        ["Testing", "15%", "15%", "→ Unchanged"],
        ["Documentation", "35%", "50%", "↗ ▲ +15 pts (Full code line traceability report)"]
    ]

    for row_idx, row_data in enumerate(data6, start=1):
        for col_idx, cell_value in enumerate(row_data):
            table6.cell(row_idx, col_idx).paragraphs[0].text = cell_value

    style_table(table6, col_widths6, col_align6)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 7 ---
    add_heading_1("7. Remaining Blockers, Ranked")
    table7 = doc.add_table(rows=7, cols=4)
    col_widths7 = [Inches(0.6), Inches(1.2), Inches(3.4), Inches(1.3)]
    col_align7 = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]

    headers7 = ["Priority", "Owner", "Blocker Description", "Severity"]
    for i, h in enumerate(headers7):
        table7.cell(0, i).paragraphs[0].text = h

    data7 = [
        ["1", "All", "firestore.rules written in repository but NOT YET DEPLOYED to production Firebase project farmerhelp-4d27c. Must execute firebase deploy --only firestore:rules.", "Critical"],
        ["2", "Bhavesh", "Marketplace payment flow is a mock status flag (processPayment()), not integrated with real UPI/Razorpay gateway.", "High"],
        ["3", "Talha", "Hosted backend connection to local Ollama server: if production backend server cannot reach localhost:11434, plan generation falls back to isEstimated: True static template.", "Medium"],
        ["4", "Talha", "Legacy database records created prior to duration validation schema carry unvalidated stage durations; test entry '<script>alert(1)</script>' requires cleanup.", "Low"],
        ["5", "Shubham", "Language selection (LocalizationService) is updated in-memory only and not persisted to user Firestore profile across app restarts.", "Low"],
        ["6", "All", "Automated test coverage: default Flutter counter smoke test remains in widget_test.dart.", "Low"]
    ]

    for row_idx, row_data in enumerate(data7, start=1):
        for col_idx, cell_value in enumerate(row_data):
            table7.cell(row_idx, col_idx).paragraphs[0].text = cell_value

    style_table(table7, col_widths7, col_align7)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 8 ---
    add_heading_1("8. Updated Supervisor Demonstration Checklist")
    add_body("The complete end-to-end smart farming journey can now be demonstrated from the feature-recommender branch:")
    add_bullet("Language Selection → Phone OTP Login → Profile Setup → Dashboard with Live GPS Location & Agro-Weather.", bold_prefix="• Core App Startup: ")
    add_bullet("Open-Meteo & SoilGrids auto-fetch → District Dropdown → Top 3 Comparative Preset Cards (N, P, K, pH, Water & DB votes) → Recommendation Generation → Results Screen with In-App Voting Card ('Was this recommendation useful?') → Firestore Vote Recording.", bold_prefix="• Crop Advisor & Voting: ")
    add_bullet("Generate AI Growth Plan via qwen2.5:7b schema → View Stage Schedule Timeline → Tap '📷 Upload Crop Photo to Refine Plan' → Process image via Qwen 2.5 VL endpoint → View interactive Recalibration Diff Dialog → Apply updated schedule.", bold_prefix="• Growth Planner AI & Vision: ")
    add_bullet("Farmer creates listing with GPS geohash → Buyer views live Google Map → Buyer places bid → Farmer accepts bid → PDF Invoice generated.", bold_prefix="• Marketplace Flow: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 9 ---
    add_heading_1("9. Final Supervisor Assessment & Updated Verdict")
    add_body(
        "AgroSaathi has progressed from late-stage integration into a feature-complete smart farming assistant. "
        "All primary user requirements across location resolution, community preset voting, past history querying, "
        "and vision-assisted growth plan recalibration are implemented in the feature-recommender branch."
    )

    table9 = doc.add_table(rows=7, cols=2)
    col_widths9 = [Inches(2.5), Inches(4.0)]
    col_align9 = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]

    headers9 = ["Metric", "Value (1 October 2026)"]
    for i, h in enumerate(headers9):
        table9.cell(0, i).paragraphs[0].text = h

    data9 = [
        ["Current Maturity", "Advanced Integration → Refinement & Pre-Submission Polish"],
        ["Overall Estimated Development", "82%"],
        ["Mainline-Integrated Implementation", "80%"],
        ["Integration Readiness Score", "84 / 100"],
        ["Strongest Components", "Crop Growth Planner (AI + Vision), Crop Recommender (Live GPS + SoilGrids + Voting), Dashboard, Marketplace"],
        ["Biggest Remaining Action Item", "Deploying firestore.rules to production Firebase project via CLI"]
    ]

    for row_idx, row_data in enumerate(data9, start=1):
        for col_idx, cell_value in enumerate(row_data):
            table9.cell(row_idx, col_idx).paragraphs[0].text = cell_value

    style_table(table9, col_widths9, col_align9)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 10 ---
    add_heading_1("10. What Must Be Completed Before Final Submission")
    add_bullet("Execute 'firebase deploy --only firestore:rules' to activate database security rules in production.", bold_prefix="1. Deploy Firestore Security Rules: ")
    add_bullet("Perform physical device camera field testing for crop photo upload and vision recalibration.", bold_prefix="2. On-Device Vision Testing: ")
    add_bullet("Ensure production backend server points to a publicly accessible or hosted Ollama instance.", bold_prefix="3. Production Backend Ollama Setup: ")
    add_bullet("Delete legacy test entries (such as '<script>alert(1)</script>') from Firestore collections.", bold_prefix="4. Database Sanity Cleanup: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_end = doc.add_paragraph()
    r_end = p_end.add_run("End of updated report.")
    r_end.font.italic = True
    r_end.font.color.rgb = COLOR_MUTED

    doc.save("AgroSaathi_Progress_Report_Updated_1Oct2026.docx")
    print("Report generated successfully as AgroSaathi_Progress_Report_Updated_1Oct2026.docx")

if __name__ == "__main__":
    create_report()
